"""Checker for the S8 figures (loopbudget). Independent of the plotting scripts: it recomputes every CSV value from
the frozen S4 files, checks the CSV source hashes, checks that every PDF embeds only TrueType fonts (each /FontFile2
stream starts with 00 01 00 00 and fontTools finds a 'glyf' table; no Type 1, CFF or Type 3 fonts), lists each PNG's
size and dpi, checks each pptx (one slide, notes = caption), and writes a timestamped log.

Usage: python check_figures.py [--dir FIGURE_DIR]   (default: the directory of this file)
Exit 0 if every check passes, 1 otherwise."""
import argparse
import csv
import hashlib
import io
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
S4 = ROOT / "research" / "s4"
SRC = {
    "judge_final": S4 / "judge_final_20261011T054347Z.json",
    "judge_first": S4 / "judge_20261011T045319Z.json",
    "selection": S4 / "selection.json",
    "n2_input": S4 / "n2_input.json",
    "profile_summary": S4 / "profile_registry_summary.json",
    "n1_input": S4 / "n1_input_topup.json",  # S9 revision: pins the confirmation run files (fig1_design_v2)
}
FIGS = ["fig1_design", "fig2_feasible_counts", "fig3_a1_cells", "fig4_a2_accounting",
        "fig1_design_v2", "fig4_a2_accounting_v2"]
# Numbering follows the order of first citation in the manuscript (renumbered 2026-10-11).
FIG_NUMBER = {"fig1_design": 1, "fig2_feasible_counts": 2, "fig3_a1_cells": 3, "fig4_a2_accounting": 4,
              "fig1_design_v2": 1, "fig4_a2_accounting_v2": 4}
A1_SENTENCE = "Bars run from the mean to the one-sided 95% upper bound."
A1_ABSENT = "not significance markers"
TIED = ("shared", "two_groups", "reversible_shared")
UNTIED = ("untied", "reversible_untied", "wide_single")
BASE = ("param_count_with_state_multiplier", "persistent_weights_gradients_optimizer", "static_four_component_sum")
TASK_ORDER = ["HAR", "KWS", "CIFAR10"]

LOG = []
FAILS = []


def log(msg):
    LOG.append(msg)
    print(msg)


def fail(msg):
    FAILS.append(msg)
    log("FAIL " + msg)


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def J(name):
    return json.loads(SRC[name].read_text())


def read_csv(path):
    header, rows = [], []
    with open(path, newline="") as fh:
        for line in fh:
            if line.startswith("# "):
                header.append(line[2:].rstrip("\n"))
            else:
                break
    with open(path, newline="") as fh:
        body = [ln for ln in fh if not ln.startswith("# ")]
    rows = list(csv.reader(body))
    return header, rows[0], rows[1:]


def same(a, b):
    """CSV cell a (string) equals the recomputed value b: exact numeric equality for numbers, else string equality."""
    if isinstance(b, bool) or b is None:
        return a == str(b)
    if isinstance(b, (int, float)):
        try:
            return float(a) == float(b)
        except ValueError:
            return False
    return a == str(b)


def compare(name, got_cols, got_rows, exp_cols, exp_rows):
    if got_cols != exp_cols:
        fail(f"{name}: columns {got_cols} != {exp_cols}")
        return
    if len(got_rows) != len(exp_rows):
        fail(f"{name}: {len(got_rows)} rows, expected {len(exp_rows)}")
        return
    bad = 0
    for i, (g, e) in enumerate(zip(got_rows, exp_rows)):
        for c, a, b in zip(exp_cols, g, e):
            if not same(a, b):
                bad += 1
                if bad <= 10:
                    fail(f"{name}: row {i + 1} column {c}: csv {a!r} != recomputed {b!r}")
    if bad == 0:
        log(f"ok   {name}: {len(exp_rows)} rows x {len(exp_cols)} columns equal the recomputation")


def check_sources(name, header, needed):
    for k in needed:
        line = f"source: {SRC[k].relative_to(ROOT)} sha256 {sha(SRC[k])}"
        if line not in header:
            fail(f"{name}: header lacks the current source line for {k}")
    log(f"ok   {name}: source hashes checked ({', '.join(needed)})")


def exp_design():
    n2, sel, prof, jf, j1 = J("n2_input"), J("selection"), J("profile_summary"), J("judge_final"), J("judge_first")
    tasks = [t for t in TASK_ORDER if any(c["config_id"].split("|")[0] == t for c in n2["configs"])]
    per = [sum(1 for c in n2["configs"] if c["config_id"].split("|")[0] == t) for t in tasks]
    assert len(set(per)) == 1
    tot = {cl: {f["n_total"] for f in sel["feasibility"] if f["class"] == cl} for cl in ("tied", "untied")}
    mb = sorted({m["microbatch"] for p in prof["profiles"].values() for m in p["measurements"]})
    ck = sorted({m["checkpointing"] for p in prof["profiles"].values() for m in p["measurements"]})
    seeds = sorted({int(re.search(r"__s(\d+)/result\.json$", cell[k]["result"]).group(1))
                    for cell in sel["cells"] for k in ("tied", "untied") if cell[k]})
    first, final = j1["n1"]["seed_set"], jf["n1"]["seed_set"]
    fl = jf["n1"]["floors"]
    j = ";".join
    rows = [["n_tasks", len(tasks)], ["tasks", j(tasks)], ["configs_per_task", per[0]],
            ["configs_total", len(n2["configs"])], ["tied_per_task", min(tot["tied"])],
            ["untied_per_task", min(tot["untied"])], ["caps_kib", j(map(str, sel["caps_kib"]))],
            ["control_cap", sel["control_cap"]], ["microbatch_min", mb[0]], ["microbatch_max", mb[-1]],
            ["checkpointing_values", j(map(str, ck))], ["selection_seeds", j(map(str, seeds))],
            ["confirmation_seeds", j(map(str, first))], ["topup_seeds", j(str(s) for s in final if s not in first)],
            ["holm_family_size", jf["n1"]["holm_family_size"]], ["floors", j(f"{t}={fl[t]!r}" for t in fl)]]
    assert len(tot["tied"]) == 1 and len(tot["untied"]) == 1
    return ["quantity", "value"], rows, ["n2_input", "selection", "profile_summary", "judge_final", "judge_first"]


def exp_a1_cells():
    jf = J("judge_final")
    cells = jf["n1"]["cells"]
    caps = sorted({c["cap_kib"] for c in cells if c["finite_cap"]})
    rows = []
    for t in TASK_ORDER:
        for c in cells:
            if c["task"] != t:
                continue
            x = caps.index(c["cap_kib"]) if c["finite_cap"] else 8.0
            rows.append([t, c["cap_kib"], x, c["mean"], c["upper_bound_95"], jf["n1"]["floors"][t], c["E_plus"],
                         c["E_minus"], c["zero_sign"], c["in_holm_family"]])
            if c["floor"] != jf["n1"]["floors"][t]:
                fail(f"fig3_a1_cells: floor mismatch in the judge for {t} {c['cap_kib']}")
    cols = ["task", "cap_kib", "x_position", "mean_delta_pp", "upper_bound_95_pp", "floor_pp", "E_plus", "E_minus",
            "zero_sign", "in_holm_family"]
    return cols, rows, ["judge_final"]


def cls(cid):
    s = cid.split("|")[1]
    return "tied" if s in TIED else "untied" if s in UNTIED else None


def exp_a2_accounting():
    n2, jf = J("n2_input"), J("judge_final")
    rows = []
    for t in TASK_ORDER:
        for cl in ("tied", "untied"):
            for c in n2["configs"]:
                if c["config_id"].split("|")[0] == t and cls(c["config_id"]) == cl:
                    rows.append(["a", c["config_id"], t, cl, c["measured_min_peak_bytes"],
                                 c["predicted_min_peak_bytes"], "", "", ""])
    right = {k: 0 for k in ("accounting", *BASE)}
    n = 0
    for c in n2["configs"]:
        m = c["measured_min_peak_bytes"]
        for cap in n2["caps_kib"]:
            B = cap * 1024
            if abs(m - B) <= 0.20 * B:
                n += 1
                fit = m <= B
                right["accounting"] += (c["predicted_min_peak_bytes"] <= B) == fit
                for b in BASE:
                    right[b] += (c["baselines"][b] <= B) == fit
    for k in ("accounting", *BASE):
        rows.append(["b", k, "", "", "", "", right[k], n, right[k] / n])
    # agreement with the frozen judge
    if not (n == jf["n2_1"]["n_boundary_pairs"] and right["accounting"] == jf["n2_1"]["n_boundary_right"]
            and right["accounting"] / n == jf["n2_1"]["boundary_fit_accuracy"]):
        fail("fig4_a2_accounting: accounting boundary accuracy differs from judge n2_1")
    for comp in jf["n2_2"]["comparisons"]:
        if right[comp["baseline"]] / n != comp["baseline_boundary_fit_accuracy"]:
            fail(f"fig4_a2_accounting: {comp['baseline']} accuracy differs from judge n2_2")
    exact = sum(c["measured_min_peak_bytes"] == c["predicted_min_peak_bytes"] for c in n2["configs"])
    log(f"info fig4_a2_accounting: boundary pairs {n}; accounting {right['accounting']}/{n}; "
        + "; ".join(f"{b} {right[b]}/{n} = {right[b] / n:.4f}" for b in BASE) + f"; exact {exact}/{len(n2['configs'])}")
    cols = ["panel", "item", "task", "class", "measured_min_peak_bytes", "predicted_min_peak_bytes", "n_right",
            "n_pairs", "fit_accuracy"]
    return cols, rows, ["n2_input", "judge_final"]


def run_knobs(pins, mode):
    """Training (microbatch, checkpointing) values over the pinned result.json files of one mode; hashes checked."""
    seen, n = set(), 0
    for rel, digest in pins.items():
        if rel.endswith("/result.json"):
            raw = (ROOT / rel).read_bytes()
            if hashlib.sha256(raw).hexdigest() != digest:
                fail(f"fig1_design_v2: {rel} differs from its pinned sha256")
            r = json.loads(raw)
            if r["mode"] != mode:
                fail(f"fig1_design_v2: {rel} mode {r['mode']} != {mode}")
            seen.add((r["microbatch"], r["checkpointing"]))
            n += 1
    if len(seen) != 1:
        fail(f"fig1_design_v2: {mode} runs have training knobs {sorted(seen)}")
    return sorted(seen)[0], n


def exp_design_v2():
    """S9 revision: the v1 quantities plus the training microbatch / checkpointing of the selection and
    confirmation runs, read from the result files pinned by selection.json and by n1_input_topup.json
    (itself pinned by judge_final input_sha256.n1)."""
    cols, rows, srcs = exp_design()
    sel, jf, n1 = J("selection"), J("judge_final"), J("n1_input")
    if sha(SRC["n1_input"]) != jf["input_sha256"]["n1"]:
        fail("fig1_design_v2: n1_input_topup.json is not the file the final judge pins")
    (mb_s, ck_s), n_s = run_knobs(sel["inputs_sha256"], "select")
    (mb_c, ck_c), n_c = run_knobs(n1["_provenance"]["inputs_sha256"], "confirm")
    if n_s != len(J("n2_input")["configs"]) or n_c != n1["_provenance"]["n_confirm_runs"]:
        fail(f"fig1_design_v2: {n_s} selection and {n_c} confirmation run files")
    rows = rows + [["train_microbatch_select", mb_s], ["train_checkpointing_select", ck_s],
                   ["n_select_runs", n_s], ["train_microbatch_confirm", mb_c],
                   ["train_checkpointing_confirm", ck_c], ["n_confirm_runs", n_c]]
    return cols, rows, srcs + ["n1_input"]


def exp_a2_accounting_v2():
    """S9 revision of Figure 4: estimate/measured ratios (with the deterministic jitter), class medians, minima and
    maxima; false fits and false rejections on the boundary pairs; checked against judge N2.1 and N2.2."""
    n2, jf = J("n2_input"), J("judge_final")
    phi = 0.6180339887498949
    rows = []
    for bi, b in enumerate(BASE):
        for ci, cl in enumerate(("tied", "untied")):
            yc = -(bi * 2.6 + ci * 1.0)
            sub = [c for c in n2["configs"] if cls(c["config_id"]) == cl]
            rs = []
            for k, c in enumerate(sub):
                r = c["baselines"][b] / c["measured_min_peak_bytes"]
                rs.append(r)
                rows.append(["a", "point", b, cl, c["config_id"], c["baselines"][b], c["measured_min_peak_bytes"], r,
                             yc + 0.6 * ((((k + 1) * phi) % 1.0) - 0.5), "", ""])
            s_ = sorted(rs)
            m = len(s_)
            med = s_[m // 2] if m % 2 else (s_[m // 2 - 1] + s_[m // 2]) / 2
            rows += [["a", "median", b, cl, "", "", "", med, yc, "", m], ["a", "min", b, cl, "", "", "", s_[0], "", "", m],
                     ["a", "max", b, cl, "", "", "", s_[-1], "", "", m]]
            if not s_[-1] < 1.0:
                fail(f"fig4_a2_accounting_v2: {b} {cl} has a ratio >= 1")
    exact = sum(c["measured_min_peak_bytes"] == c["predicted_min_peak_bytes"] for c in n2["configs"])
    rows.append(["a", "accounting_exact", "accounting", "", "", "", "", 1.0, "", exact, len(n2["configs"])])
    meth = (*BASE, "accounting")
    false_fit = dict.fromkeys(meth, 0)
    false_rej = dict.fromkeys(meth, 0)
    n_fit = n_not = 0
    for c in n2["configs"]:
        m = c["measured_min_peak_bytes"]
        est = dict(c["baselines"], accounting=c["predicted_min_peak_bytes"])
        for cap in n2["caps_kib"]:
            B = cap * 1024
            if abs(m - B) <= 0.20 * B:
                if m <= B:
                    n_fit += 1
                    for k in meth:
                        false_rej[k] += est[k] > B
                else:
                    n_not += 1
                    for k in meth:
                        false_fit[k] += est[k] <= B
    rows += [["b", "false_fit", k, "", "", "", "", "", "", false_fit[k], n_not] for k in meth]
    rows += [["b", "false_reject", k, "", "", "", "", "", "", false_rej[k], n_fit] for k in meth]
    # agreement with the frozen judge and with the counts named by the S9 card
    if not (n_fit + n_not == jf["n2_1"]["n_boundary_pairs"] == 928 and n_fit == 500 and n_not == 428):
        fail(f"fig4_a2_accounting_v2: boundary pairs {n_fit} fit + {n_not} not fit")
    if not (exact == jf["n2_1"]["n_configs"] == 744 and jf["n2_1"]["max_rel_err"] == 0.0
            and false_fit["accounting"] == false_rej["accounting"] == 0 and jf["n2_1"]["n_boundary_right"] == 928):
        fail("fig4_a2_accounting_v2: accounting disagrees with judge n2_1")
    for comp in jf["n2_2"]["comparisons"]:
        k = comp["baseline"]
        if not (false_fit[k] + false_rej[k] == comp["b"] and comp["c"] == 0 and false_rej[k] == 0):
            fail(f"fig4_a2_accounting_v2: {k} disagrees with judge n2_2 (b {comp['b']}, c {comp['c']})")
    log(f"info fig4_a2_accounting_v2: pairs {n_fit + n_not} (fit {n_fit}, not fit {n_not}); false fits "
        + ", ".join(f"{k} {false_fit[k]}" for k in meth) + "; false rejections "
        + ", ".join(f"{k} {false_rej[k]}" for k in meth) + f"; exact {exact}/{len(n2['configs'])}")
    cols = ["panel", "row", "method", "class", "config_id", "estimate_bytes", "measured_min_peak_bytes", "value",
            "y_position", "count", "n"]
    return cols, rows, ["n2_input", "judge_final"]


def exp_feasible_counts():
    sel = J("selection")
    order = sel["caps_kib"] + ["unbounded"]
    f = {(x["task"], x["class"], x["cap_kib"]): x for x in sel["feasibility"]}
    rows = []
    for t in TASK_ORDER:
        for cl in ("tied", "untied"):
            for cap in order:
                x = f[(t, cl, cap)]
                rows.append([t, cl, cap, order.index(cap) if cap != "unbounded" else 8.0, x["n_feasible"],
                             x["n_total"]])
    return ["task", "class", "cap_kib", "x_position", "n_feasible", "n_total"], rows, ["selection"]


def check_pdf(path):
    import pymupdf as fitz
    from fontTools.ttLib import TTFont

    doc = fitz.open(path)
    n_ff2 = 0
    names = set()
    for x in range(1, doc.xref_length()):
        keys = doc.xref_get_keys(x)
        if "Subtype" in keys and doc.xref_get_key(x, "Subtype")[1] == "/Type3":
            fail(f"{path.name}: Type3 font at xref {x}")
        if "BaseFont" in keys:
            names.add(doc.xref_get_key(x, "BaseFont")[1])
        if doc.xref_get_key(x, "Type")[1] == "/FontDescriptor":
            if "FontFile" in keys or "FontFile3" in keys:
                fail(f"{path.name}: non-TrueType font program at xref {x}")
            if "FontFile2" not in keys:
                fail(f"{path.name}: font descriptor {x} without embedded FontFile2")
                continue
            ref = int(doc.xref_get_key(x, "FontFile2")[1].split()[0])
            data = doc.xref_stream(ref)
            if data[:4] != b"\x00\x01\x00\x00":
                fail(f"{path.name}: FontFile2 {ref} starts with {data[:4].hex()}")
            if "glyf" not in TTFont(io.BytesIO(data)):
                fail(f"{path.name}: FontFile2 {ref} has no glyf table")
            n_ff2 += 1
    if n_ff2 == 0:
        fail(f"{path.name}: no FontFile2 stream")
    if any("DejaVu" not in nm for nm in names):
        fail(f"{path.name}: non-DejaVu font names {sorted(names)}")
    w, h = doc[0].rect.width, doc[0].rect.height
    log(f"ok   {path.name}: {len(doc)} page, {w / 72 * 2.54:.2f} x {h / 72 * 2.54:.2f} cm; {n_ff2} FontFile2 "
        f"stream(s), 00 01 00 00 + glyf; fonts {sorted(names)}")


def check_png(path):
    from PIL import Image

    with Image.open(path) as im:
        w, h = im.size
        dpi = im.info.get("dpi")
    if not dpi or abs(dpi[0] - 300) > 0.5 or abs(dpi[1] - 300) > 0.5:
        fail(f"{path.name}: dpi {dpi}")
    log(f"ok   {path.name}: {w} x {h} px, dpi {tuple(round(d, 2) for d in dpi) if dpi else None}, "
        f"{w / 300 * 2.54:.2f} x {h / 300 * 2.54:.2f} cm at 300 dpi")


def check_pptx(path, caption):
    from pptx import Presentation

    prs = Presentation(path)
    slides = list(prs.slides)
    if len(slides) != 1:
        fail(f"{path.name}: {len(slides)} slides")
        return
    notes = slides[0].notes_slide.notes_text_frame.text if slides[0].has_notes_slide else ""
    if notes != caption:
        fail(f"{path.name}: notes differ from the CSV caption")
    pics = [s for s in slides[0].shapes if s.shape_type == 13]
    if len(pics) != 1 or pics[0].width != prs.slide_width:
        fail(f"{path.name}: expected one picture at slide width")
    log(f"ok   {path.name}: 1 slide, picture at slide width, notes = caption")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(Path(__file__).resolve().parent))
    d = Path(ap.parse_args().dir).resolve()
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    log(f"check_figures.py run {stamp} on {d}")
    for k, p in SRC.items():
        log(f"src  {k}: {p.relative_to(ROOT)} sha256 {sha(p)}")
    expect = {"fig1_design": exp_design, "fig2_feasible_counts": exp_feasible_counts,
              "fig3_a1_cells": exp_a1_cells, "fig4_a2_accounting": exp_a2_accounting,
              "fig1_design_v2": exp_design_v2, "fig4_a2_accounting_v2": exp_a2_accounting_v2}
    for name in FIGS:
        for ext in ("pdf", "png", "csv", "pptx"):
            if not (d / f"{name}.{ext}").is_file():
                fail(f"{name}.{ext} missing")
        if not (d / f"{name}.csv").is_file():
            continue
        header, cols, rows = read_csv(d / f"{name}.csv")
        ecols, erows, srcs = expect[name]()
        compare(name, cols, rows, ecols, erows)
        check_sources(name, header, srcs)
        caption = header[0] if header else ""
        if not caption.startswith(f"Figure {FIG_NUMBER[name]}. "):
            fail(f"{name}: caption does not start with 'Figure {FIG_NUMBER[name]}. '")
        if name == "fig3_a1_cells" and (A1_SENTENCE not in caption or A1_ABSENT in caption):
            fail("fig3_a1_cells: caption must contain the bar sentence and not the retired markers sentence")
        if (d / f"{name}.pdf").is_file():
            check_pdf(d / f"{name}.pdf")
        if (d / f"{name}.png").is_file():
            check_png(d / f"{name}.png")
        if (d / f"{name}.pptx").is_file():
            check_pptx(d / f"{name}.pptx", caption)
    status = "PASS" if not FAILS else f"FAIL ({len(FAILS)} failures)"
    log(f"RESULT {status}")
    (d / "logs").mkdir(exist_ok=True)
    out = d / "logs" / f"check_figures_{stamp}.log"
    out.write_text("\n".join(LOG) + "\n")
    print(f"log written to {out}")
    return 0 if not FAILS else 1


if __name__ == "__main__":
    sys.exit(main())
