"""Shared helpers for the S8 figures (loopbudget). Reads only the frozen S4 files listed in FROZEN.

Style: matplotlib's bundled DejaVu Sans (TrueType), pdf.fonttype 42, Okabe-Ito colours, sizes in cm at print size.
Outputs per figure: <name>.pdf (vector), <name>.png (300 dpi), <name>.csv (plotted numbers), <name>.pptx (one slide).
"""
import csv
import hashlib
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[4]
OUT = ROOT / "result" / "figure" / "s8"
S4 = ROOT / "research" / "s4"
FROZEN = {
    "judge_final": S4 / "judge_final_20261011T054347Z.json",
    "judge_first": S4 / "judge_20261011T045319Z.json",
    "selection": S4 / "selection.json",
    "n2_input": S4 / "n2_input.json",
    "profile_summary": S4 / "profile_registry_summary.json",
}

# Okabe-Ito
OI = {"black": "#000000", "orange": "#E69F00", "sky": "#56B4E9", "green": "#009E73", "yellow": "#F0E442",
      "blue": "#0072B2", "vermillion": "#D55E00", "purple": "#CC79A7", "grey": "#999999"}
TASKS = ["HAR", "KWS", "CIFAR10"]
TASK_NAME = {"HAR": "UCI-HAR", "KWS": "Speech Commands", "CIFAR10": "CIFAR-10"}
TASK_COLOR = {"HAR": OI["orange"], "KWS": OI["green"], "CIFAR10": OI["blue"]}
CLASS_NAME = {"tied": "depth-tied", "untied": "distinct-parameter"}
CLASS_COLOR = {"tied": OI["blue"], "untied": OI["vermillion"]}
TIED_STRUCTURES = ("shared", "two_groups", "reversible_shared")
UNTIED_STRUCTURES = ("untied", "reversible_untied", "wide_single")
BASELINES = ("param_count_with_state_multiplier", "persistent_weights_gradients_optimizer",
             "static_four_component_sum")
BASELINE_NAME = {"param_count_with_state_multiplier": "parameter count × state slots",
                 "persistent_weights_gradients_optimizer": "weights + gradients + optimizer state",
                 "static_four_component_sum": "static four-component sum"}
CM = 1 / 2.54


def setup_style():
    matplotlib.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 7.5,
        "axes.titlesize": 8,
        "axes.labelsize": 7.5,
        "xtick.labelsize": 7,
        "ytick.labelsize": 7,
        "legend.fontsize": 7,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
        "mathtext.fontset": "dejavusans",
        "axes.linewidth": 0.6,
        "xtick.major.width": 0.6,
        "ytick.major.width": 0.6,
        "lines.linewidth": 0.9,
        "savefig.dpi": 300,
    })


def load(name):
    with open(FROZEN[name]) as fh:
        return json.load(fh)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def class_of(config_id):
    s = config_id.split("|")[1]
    if s in TIED_STRUCTURES:
        return "tied"
    if s in UNTIED_STRUCTURES:
        return "untied"
    raise ValueError(f"unknown structure in {config_id}")


def write_csv(name, header_lines, columns, rows, sources):
    """CSV with '# ' comment header (caption, sources with sha256), then a column row and data rows."""
    path = OUT / f"{name}.csv"
    with open(path, "w", newline="") as fh:
        for line in header_lines:
            fh.write(f"# {line}\n")
        for s in sources:
            fh.write(f"# source: {FROZEN[s].relative_to(ROOT)} sha256 {sha256(FROZEN[s])}\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(columns)
        for r in rows:
            w.writerow(r)
    return path


def save(fig, name):
    fig.savefig(OUT / f"{name}.pdf", metadata={"CreationDate": None})
    fig.savefig(OUT / f"{name}.png", dpi=300)
    plt.close(fig)


def write_pptx(name, title, caption):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from PIL import Image

    png = OUT / f"{name}.png"
    with Image.open(png) as im:
        w_px, h_px = im.size
    slide_w = Emu(12192000)  # 13.333 in (16:9 width)
    title_h = Emu(914400)  # 1.0 in title band
    pic_h = Emu(int(slide_w * h_px / w_px))
    prs = Presentation()
    prs.slide_width = slide_w
    prs.slide_height = Emu(title_h + pic_h)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(Emu(0), Emu(0), slide_w, title_h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    slide.shapes.add_picture(str(png), Emu(0), title_h, width=slide_w, height=pic_h)
    slide.notes_slide.notes_text_frame.text = caption
    prs.save(OUT / f"{name}.pptx")
